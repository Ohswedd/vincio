"""Tests for the richer document inputs: new formats, OCR, transcripts, forms."""

from __future__ import annotations

import struct
import zipfile
import zlib
from pathlib import Path

import pytest

from vincio.documents import (
    HeuristicFormExtractor,
    MockTranscriber,
    figure_evidence,
    form_fields_to_evidence,
    load_document,
    load_media,
    register_loader,
)
from vincio.documents.parsers import parse_html, structure_data
from vincio.documents.registry import default_parser_registry


def _png(path: Path) -> Path:
    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = (bytes([0]) + bytes([10, 20, 30])) * 1
    data = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(data)
    return path


# -- new format loaders ------------------------------------------------------


class TestNewFormats:
    def test_pptx(self, tmp_path):
        p = tmp_path / "deck.pptx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("ppt/slides/slide1.xml", "<p><a:t>Title slide</a:t><a:t>Point one</a:t></p>")
            z.writestr("ppt/slides/slide2.xml", "<p><a:t>Second slide</a:t></p>")
        doc = load_document(p)
        assert "Title slide" in doc.text and len(doc.sections) == 2
        assert doc.metadata["extractor"] == "pptx-ooxml"

    def test_epub_with_spine(self, tmp_path):
        p = tmp_path / "book.epub"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("META-INF/container.xml",
                       '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/>'
                       "</rootfiles></container>")
            z.writestr("OEBPS/content.opf",
                       '<package><manifest><item id="c1" href="ch1.xhtml"/>'
                       '<item id="c2" href="ch2.xhtml"/></manifest>'
                       '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
            z.writestr("OEBPS/ch1.xhtml", "<html><body><p>Chapter one body.</p></body></html>")
            z.writestr("OEBPS/ch2.xhtml", "<html><body><p>Chapter two body.</p></body></html>")
        doc = load_document(p)
        assert "Chapter one" in doc.text and "Chapter two" in doc.text
        assert len(doc.sections) == 2

    def test_rtf(self, tmp_path):
        p = tmp_path / "note.rtf"
        p.write_text(r"{\rtf1\ansi Hello \b bold\b0 world.\par Next line.}")
        doc = load_document(p)
        assert "Hello" in doc.text and "Next line" in doc.text and doc.media_type == "application/rtf"

    def test_rtf_control_words_not_treated_as_unicode(self, tmp_path):
        # \uc1, \ulN, \upN start with 'u' but are NOT \u Unicode escapes; only
        # the exact \uNNNN escape should emit a character.
        p = tmp_path / "u.rtf"
        p.write_text(r"{\rtf1\ansi\uc1\ul0 Caf\u233  test.\ul1 done}")
        doc = load_document(p)
        assert "\x00" not in doc.text and "\x01" not in doc.text, repr(doc.text)
        assert "Café" in doc.text and "test" in doc.text and "done" in doc.text

    def test_epub_href_before_id(self, tmp_path):
        # OPF allows attributes in any order; href-before-id must still resolve.
        p = tmp_path / "rev.epub"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("META-INF/container.xml",
                       '<c><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></c>')
            z.writestr("OEBPS/content.opf",
                       '<package><manifest><item href="ch1.xhtml" id="c1"/>'
                       '<item href="ch2.xhtml" id="c2"/></manifest>'
                       '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
            z.writestr("OEBPS/ch1.xhtml", "<html><body><p>Chapter one body.</p></body></html>")
            z.writestr("OEBPS/ch2.xhtml", "<html><body><p>Chapter two body.</p></body></html>")
        doc = load_document(p)
        assert "Chapter one" in doc.text and "Chapter two" in doc.text and len(doc.sections) == 2

    def test_zip_decompression_guard(self, tmp_path):
        from vincio.core.errors import LoaderError
        from vincio.documents.formats import _read_zip_entry

        p = tmp_path / "z.pptx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("ppt/slides/slide1.xml", "<p><a:t>" + "x" * 1000 + "</a:t></p>")
        with zipfile.ZipFile(p) as archive:
            with pytest.raises(LoaderError):
                _read_zip_entry(archive, "ppt/slides/slide1.xml", max_bytes=64)

    def test_odt(self, tmp_path):
        p = tmp_path / "doc.odt"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("content.xml", "<d><text:h>Heading</text:h><text:p>Para body.</text:p></d>")
        doc = load_document(p)
        assert "Heading" in doc.text and "Para body" in doc.text

    def test_mbox(self, tmp_path):
        p = tmp_path / "thread.mbox"
        p.write_text(
            "From a@b.com Mon Jan 1 00:00:00 2024\n"
            "From: a@b.com\nSubject: Hello\n\nFirst message body.\n\n"
            "From c@d.com Mon Jan 1 00:01:00 2024\n"
            "From: c@d.com\nSubject: Re: Hello\n\nReply body.\n\n"
        )
        doc = load_document(p)
        assert doc.metadata["message_count"] == 2 and "First message body" in doc.text

    def test_parquet(self, tmp_path):
        pq = pytest.importorskip("pyarrow.parquet")
        import pyarrow as pa

        table = pa.table({"a": [1, 2], "b": ["x", "y"]})
        p = tmp_path / "data.parquet"
        pq.write_table(table, str(p))
        doc = load_document(p)
        assert doc.tables and doc.tables[0]["columns"] == ["a", "b"]


# -- parser registry ---------------------------------------------------------


class TestParserRegistry:
    def test_register_custom_loader(self, tmp_path):
        from vincio.core.types import Document

        @register_loader(".widget")
        def _load_widget(path, **_):
            return Document(text=f"widget:{path.read_text()}", media_type="application/x-widget")

        try:
            p = tmp_path / "thing.widget"
            p.write_text("payload")
            assert default_parser_registry().supports(".widget")
            doc = load_document(p)
            assert doc.text == "widget:payload"
        finally:
            default_parser_registry()._loaders.pop(".widget", None)


# -- HTML / structured data --------------------------------------------------


class TestStructuredParsing:
    def test_parse_html_tables(self):
        html = ("<html><head><title>Rpt</title></head><body><h1>Q2</h1><p>Up.</p>"
                "<table><tr><th>M</th><th>V</th></tr><tr><td>R</td><td>9</td></tr></table></body></html>")
        title, text, sections, tables = parse_html(html)
        assert title == "Rpt" and tables[0].columns == ["M", "V"] and tables[0].rows == [["R", "9"]]

    def test_structure_list_of_dicts(self):
        text, sections, tables = structure_data([{"a": 1, "b": 2}, {"a": 3, "b": 4}], title="rows")
        assert tables and tables[0].rows == [["1", "2"], ["3", "4"]]

    def test_structure_mapping(self):
        text, sections, tables = structure_data({"name": "Acme", "rows": [{"k": "v"}]})
        assert any(s.title == "name" for s in sections)
        assert any(t.title == "rows" for t in tables)

    def test_json_loader_structures(self, tmp_path):
        p = tmp_path / "d.json"
        p.write_text('[{"x":1,"y":2},{"x":3,"y":4}]')
        doc = load_document(p)
        assert doc.tables and doc.tables[0]["columns"] == ["x", "y"]


# -- audio transcript ingestion ----------------------------------------------


class TestAudioIngestion:
    def test_load_media_segments(self, tmp_path):
        wav = tmp_path / "a.wav"
        wav.write_bytes(b"RIFF\x00\x00\x00\x00WAVE")
        doc = load_media(wav, transcriber=MockTranscriber(diarize=True))
        assert doc.media_type == "audio/transcript"
        assert doc.sections and doc.sections[0]["speaker"] == "Speaker 1"
        assert doc.metadata["segment_count"] == len(doc.sections)

    async def test_provider_audio_transcriber(self, tmp_path):
        from vincio.documents.audio import ProviderAudioTranscriber
        from vincio.providers import MockProvider

        wav = tmp_path / "a.wav"
        wav.write_bytes(b"RIFF\x00\x00\x00\x00WAVE")
        transcriber = ProviderAudioTranscriber(MockProvider(default_text="hello there"), model="mock-1")
        transcript = await transcriber.transcribe(wav)
        assert "hello there" in transcript.text and transcript.segments


# -- forms / KYC -------------------------------------------------------------


class TestForms:
    def test_heuristic_extraction(self):
        fields = HeuristicFormExtractor().extract(
            "Name: Jane Doe\nDate of Birth: 1990-01-01\nInvoice Number: INV-42\nTotal: $1,200\n"
            "Random Label: some value"
        )
        names = {f.name for f in fields}
        assert {"name", "date_of_birth", "invoice_number", "total"} <= names
        canonical = next(f for f in fields if f.name == "invoice_number")
        assert canonical.confidence > 0.8 and canonical.value == "INV-42"

    def test_fields_to_evidence(self):
        fields = HeuristicFormExtractor().extract("Total: $9.99\nName: A B")
        evidence = form_fields_to_evidence(fields, source_id="KYC1", source_uri="file://x")
        assert len(evidence) == len(fields)
        assert evidence[0].source_id == "KYC1"
        assert all(":" in e.text for e in evidence)

    async def test_extract_fields_from_file(self, tmp_path):
        p = tmp_path / "form.txt"
        p.write_text("Name: Jane\nTotal: $5")
        fields = await HeuristicFormExtractor().extract_fields(p)
        assert {f.name for f in fields} == {"name", "total"}


class TestCloudDocumentAI:
    """The cloud adapters are dependency-injected (you pass the SDK client); the
    response parsers are pure and tested offline against synthetic responses."""

    def test_textract_parse(self):
        from vincio.documents.forms import TextractDocumentAI

        response = {
            "Blocks": [
                {"Id": "k1", "BlockType": "KEY_VALUE_SET", "EntityTypes": ["KEY"],
                 "Confidence": 99.0, "Page": 1,
                 "Geometry": {"BoundingBox": {"Left": 0.1, "Top": 0.2, "Width": 0.3, "Height": 0.05}},
                 "Relationships": [{"Type": "VALUE", "Ids": ["v1"]}, {"Type": "CHILD", "Ids": ["w1"]}]},
                {"Id": "v1", "BlockType": "KEY_VALUE_SET", "EntityTypes": ["VALUE"],
                 "Relationships": [{"Type": "CHILD", "Ids": ["w2"]}]},
                {"Id": "w1", "BlockType": "WORD", "Text": "Name"},
                {"Id": "w2", "BlockType": "WORD", "Text": "Jane Doe"},
            ]
        }
        fields = TextractDocumentAI.parse(response)
        assert len(fields) == 1
        f = fields[0]
        assert f.name == "name" and f.value == "Jane Doe" and f.page == 1
        assert abs(f.confidence - 0.99) < 1e-6
        assert f.bbox == (0.1, 0.2, pytest.approx(0.4), pytest.approx(0.25))

    async def test_textract_async_with_fake_client(self, tmp_path):
        from vincio.documents.forms import TextractDocumentAI

        captured = {}

        class FakeTextract:
            def analyze_document(self, **kwargs):
                captured.update(kwargs)
                return {"Blocks": [
                    {"Id": "k", "BlockType": "KEY_VALUE_SET", "EntityTypes": ["KEY"],
                     "Relationships": [{"Type": "VALUE", "Ids": ["v"]}, {"Type": "CHILD", "Ids": ["wk"]}]},
                    {"Id": "v", "BlockType": "KEY_VALUE_SET", "EntityTypes": ["VALUE"],
                     "Relationships": [{"Type": "CHILD", "Ids": ["wv"]}]},
                    {"Id": "wk", "BlockType": "WORD", "Text": "Total"},
                    {"Id": "wv", "BlockType": "WORD", "Text": "$5"},
                ]}

        doc = tmp_path / "scan.png"
        doc.write_bytes(b"\x89PNG bytes")
        fields = await TextractDocumentAI(FakeTextract()).extract_fields(doc)
        assert captured["FeatureTypes"] == ["FORMS"] and captured["Document"]["Bytes"]
        assert fields[0].name == "total" and fields[0].value == "$5"

    def test_azure_parse(self):
        from vincio.documents.forms import AzureDocumentAI

        result = {"key_value_pairs": [
            {"key": {"content": "Date of Birth",
                     "bounding_regions": [{"page_number": 1, "polygon": [0.1, 0.1, 0.4, 0.1, 0.4, 0.2, 0.1, 0.2]}]},
             "value": {"content": "1990-01-01"}, "confidence": 0.97},
            {"key": {"content": ""}, "value": {"content": "skip"}, "confidence": 0.5},
        ]}
        fields = AzureDocumentAI.parse(result)
        assert len(fields) == 1
        assert fields[0].name == "date_of_birth" and fields[0].value == "1990-01-01"
        assert fields[0].page == 1 and fields[0].bbox == (0.1, 0.1, 0.4, 0.2)

    def test_google_parse(self):
        from vincio.documents.forms import GoogleDocumentAI

        document = {
            "text": "Invoice Number: INV-9\n",
            "pages": [{"page_number": 1, "form_fields": [
                {"field_name": {"text_anchor": {"text_segments": [{"start_index": 0, "end_index": 14}]}},
                 "field_value": {"text_anchor": {"text_segments": [{"start_index": 16, "end_index": 21}]}},
                 "field_name_confidence": 0.9}]}],
        }
        fields = GoogleDocumentAI.parse(document)
        assert len(fields) == 1
        assert fields[0].name == "invoice_number" and fields[0].value == "INV-9" and fields[0].page == 1


class TestOptionalDepErrors:
    """Missing optional deps must fail with a clear, actionable message."""

    def test_pptx_render_without_lib(self):
        try:
            import pptx  # noqa: F401
        except ImportError:
            from vincio.core.errors import GenerationError
            from vincio.generation.model import DocumentModel
            from vincio.generation.render import render

            model = DocumentModel(title="T")
            model.paragraph("body")
            with pytest.raises(GenerationError, match="gen-pptx"):
                render(model, "pptx")
        else:
            pytest.skip("python-pptx installed; missing-dep path not exercisable")

    def test_parquet_loader_without_lib(self, tmp_path):
        try:
            import pyarrow  # noqa: F401
        except ImportError:
            from vincio.core.errors import LoaderError

            p = tmp_path / "x.parquet"
            p.write_bytes(b"not really parquet")
            with pytest.raises(LoaderError, match="parquet"):
                load_document(p)
        else:
            pytest.skip("pyarrow installed; missing-dep path not exercisable")

    def test_msg_loader_without_lib(self, tmp_path):
        try:
            import extract_msg  # noqa: F401
        except ImportError:
            from vincio.core.errors import LoaderError

            p = tmp_path / "x.msg"
            p.write_bytes(b"\xd0\xcf\x11\xe0 fake ole")
            with pytest.raises(LoaderError, match="extract-msg"):
                load_document(p)
        else:
            pytest.skip("extract-msg installed; missing-dep path not exercisable")


# -- figure → evidence -------------------------------------------------------


class TestFigureEvidence:
    async def test_figures_to_evidence_with_analyzer(self, tmp_path):
        from vincio.core.types import Document
        from vincio.documents.multimodal import ImageAnalyzer
        from vincio.providers import MockProvider

        crop = _png(tmp_path / "fig0.png")
        document = Document(
            text="report", metadata={"figures": [{"page": 2, "bbox": [0, 0, 100, 80]}]}
        )
        analyzer = ImageAnalyzer(MockProvider(), model="mock-1")
        evidence = figure_evidence(document, crops={0: str(crop)}, analyzer=analyzer)
        assert evidence and evidence[0].source_type == "image"
        assert evidence[0].metadata["bbox"] == [0, 0, 100, 80]
        assert evidence[0].page == 2

    async def test_figures_to_evidence_with_ocr(self, tmp_path):
        from vincio.core.types import Document

        crop = _png(tmp_path / "fig0.png")

        class _OCR:
            async def extract_text(self, image_path):
                return "Chart shows growth"

        document = Document(text="x", metadata={"figures": [{"page": 1, "bbox": [1, 2, 3, 4]}]})
        evidence = figure_evidence(document, crops={0: str(crop)}, ocr_engine=_OCR())
        assert evidence and "Chart shows growth" in evidence[0].text


# -- OCR auto-fallback in load_pdf -------------------------------------------


class TestPDFOCR:
    def test_ocr_fallback_routes_low_text_pages(self, tmp_path, monkeypatch):
        pytest.importorskip("pypdf")
        from pypdf import PdfWriter

        # A blank single-page PDF yields ~no text → OCR fallback fires.
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        pdf_path = tmp_path / "scan.pdf"
        with pdf_path.open("wb") as fh:
            writer.write(fh)

        class _OCR:
            async def extract_text(self, image_path):
                return "Recovered scanned text content"

        def fake_raster(path, idx):
            return str(_png(tmp_path / f"pg{idx}.png"))

        from vincio.documents.loaders import load_pdf

        doc = load_pdf(pdf_path, ocr_engine=_OCR(), rasterizer=fake_raster)
        assert "Recovered scanned text" in doc.text
        assert doc.metadata.get("ocr_pages") == [1]
        assert doc.sections[0]["extractor"] == "ocr"
